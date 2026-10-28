"""Metadata-gated Graph sync. Web requests consume SQLite only."""
from datetime import datetime, timezone
from pathlib import Path
import os

from filelock import FileLock, Timeout
from flask import current_app
from app import db
from app.models import DataSourceState, SyncRun
from app.services.data_sources import SOURCES, resolve, fetch_source
from app.integrations.onedrive import OneDriveClient, GraphError
from app.integrations.graph_auth import GraphAuthError


def _now():
    return datetime.now(timezone.utc)


def version(key):
    if key == "finance_workbook":
        from app.importer import IMPORTER_VERSION
        return IMPORTER_VERSION + ":mortgage-1"
    if key == "gymday_latest_zip":
        from app.gymday import IMPORT_VERSION
        return IMPORT_VERSION
    return "body-composition-1"


def _import(key, payload, force):
    if key == "finance_workbook":
        from app.importer import import_bytes
        from app.finance import mortgage_tracking
        from openpyxl import load_workbook
        from io import BytesIO
        # Validate the mortgage sheet before importing; preserve cached formula results.
        wb = load_workbook(BytesIO(payload.content), read_only=True, data_only=True)
        try:
            if "Net Worth Tracking" not in wb.sheetnames:
                raise ValueError("Required Net Worth Tracking sheet missing")
        finally:
            wb.close()
        outcome = import_bytes(payload.content, payload.name, force=force)
        # Atomic replacement of managed local cache: no ordinary chart calls to Graph.
        cache = Path(current_app.instance_path) / "finance-graph-cache.xlsx"
        cache.parent.mkdir(parents=True, exist_ok=True)
        pending = cache.with_suffix(".pending")
        try:
            pending.write_bytes(payload.content)
            os.replace(pending, cache)
        finally:
            pending.unlink(missing_ok=True)
        return outcome
    if key == "body_composition_workbook":
        from app.body_composition import import_source_bytes
        return import_source_bytes(payload.content, payload.name, force=force)
    if key == "gymday_latest_zip":
        from app.gymday import import_zip_bytes
        return import_zip_bytes(payload.content, payload.name, force=force)
    raise ValueError("Unknown data source")


def _local_fallback(key, state):
    if not current_app.config.get("LIFE_SYNC_LOCAL_FALLBACK"):
        return None
    # Once Graph has succeeded, local sync never replaces newer remote data.
    if state.last_source_type == "onedrive":
        return None
    if key == "finance_workbook":
        from app.importer import import_path
        path = current_app.config.get("SOURCE_WORKBOOK")
        if path and Path(path).is_file():
            return import_path(path)
    elif key == "body_composition_workbook":
        from app.body_composition import import_source
        return import_source()
    elif key == "gymday_latest_zip":
        from app.gymday import import_latest
        return import_latest()
    return None


def sync_one(key, force=False, trigger="manual", client=None):
    if key not in SOURCES:
        raise ValueError("Unknown source key")
    if current_app.config.get("LIFE_DATA_SOURCE") != "onedrive":
        raise ValueError("Graph sync requires LIFE_DATA_SOURCE=onedrive")
    app = current_app._get_current_object()
    Path(app.instance_path).mkdir(parents=True, exist_ok=True)
    lock = FileLock(str(Path(app.instance_path) / ("sync-" + key + ".lock")))
    try:
        lock.acquire(timeout=0)
    except Timeout:
        return {"status": "already_running", "source_key": key}
    try:
        state = DataSourceState.query.filter_by(source_key=key).first()
        if state is None:
            state = DataSourceState(source_key=key)
            db.session.add(state)
            db.session.flush()
        run = SyncRun(source_key=key, trigger=trigger)
        db.session.add(run)
        db.session.commit()
        try:
            api = client or OneDriveClient()
            item, drive_id = resolve(api, key, state if key != "gymday_latest_zip" else None)
            state.last_checked_at = _now()
            etag = item.get("eTag")
            if (not force and state.status in ("success", "unchanged") and
                state.last_source_type == "onedrive" and state.item_id == item["id"] and
                state.importer_version == version(key) and etag and state.etag == etag):
                run.status = state.status = "unchanged"
                run.finished_at = _now()
                db.session.commit()
                return {"status": "unchanged", "run_id": run.id}
            payload = fetch_source(api, key, item, drive_id)
            run.bytes_downloaded = len(payload.content)
            if (not force and state.sha256 == payload.sha256 and
                state.importer_version == version(key) and state.last_source_type == "onedrive"):
                result = {"status": "unchanged", "counts": {}}
            else:
                result = _import(key, payload, force=force)
            state.drive_id, state.item_id = payload.drive_id, payload.item_id
            state.etag, state.remote_modified_at = payload.etag, payload.modified_at
            state.sha256, state.importer_version = payload.sha256, version(key)
            state.last_source_type = "onedrive"
            state.provider = "onedrive"
            state.last_success_at = _now()
            state.last_error_code = state.last_error_summary = None
            state.status = "success" if result["status"] not in ("unchanged",) else "unchanged"
            run.status, run.finished_at = state.status, _now()
            counts = result.get("counts") or {}
            run.records_imported = sum(v for v in counts.values() if isinstance(v, int))
            db.session.commit()
            return {"status": run.status, "run_id": run.id, "counts": counts}
        except Exception as exc:
            db.session.rollback()
            state = DataSourceState.query.filter_by(source_key=key).one()
            run = db.session.get(SyncRun, run.id)
            code = exc.code if isinstance(exc, (GraphError, GraphAuthError)) else "invalid_workbook" if isinstance(exc, ValueError) else "import_failed"
            safe = {"not_found": "Configured OneDrive source not found",
                    "reconnect_required": "Microsoft sign-in required",
                    "forbidden": "Graph Files.Read permission denied",
                    "throttled": "Graph rate limit reached",
                    "unavailable": "Microsoft Graph temporarily unavailable",
                    "invalid_workbook": "Source validation failed"}
            summary = safe.get(code, "Source sync failed; inspect configuration and retry")
            state.last_checked_at = _now()
            state.last_error_code, state.last_error_summary = code, summary
            state.status = code
            run.status, run.error_code, run.error_summary = "failed", code, summary
            run.finished_at = _now()
            db.session.commit()
            try:
                fallback = _local_fallback(key, state)
                if fallback and fallback.get("status") in ("success", "completed", "unchanged"):
                    state.last_source_type = "local"
                    state.last_success_at = _now()
                    state.status = "local_fallback"
                    db.session.commit()
                    return {"status": "local_fallback", "run_id": run.id, "remote_error": code}
            except Exception:
                db.session.rollback()
            return {"status": "failed", "run_id": run.id, "error_code": code, "message": summary}
    finally:
        lock.release()


def sync_all(force=False, trigger="manual"):
    return {key: sync_one(key, force=force, trigger=trigger) for key in SOURCES}


def statuses():
    items = {s.source_key: s for s in DataSourceState.query.all()}
    return [{"source_key": key,
             "status": items[key].status if key in items else "not_synced",
             "last_checked_at": items[key].last_checked_at.isoformat() if key in items and items[key].last_checked_at else None,
             "last_success_at": items[key].last_success_at.isoformat() if key in items and items[key].last_success_at else None,
             "importer_version": items[key].importer_version if key in items else None,
             "last_source_type": items[key].last_source_type if key in items else None,
             "error_code": items[key].last_error_code if key in items else None,
             "error_summary": items[key].last_error_summary if key in items else None}
            for key in SOURCES]
