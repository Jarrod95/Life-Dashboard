"""Read-only, source-separated body composition ingestion and endpoints."""
from __future__ import annotations

import hashlib
import math
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from io import BytesIO

from flask import Blueprint, current_app, jsonify, render_template, request
from openpyxl import load_workbook

from app import db

pages = Blueprint("body_composition_pages", __name__)
api = Blueprint("body_composition_api", __name__, url_prefix="/api/v1/fitness/body-composition")
SHEETS = {"manual": "Data", "evolt": "Evolt Scans",
          "dexa": "DEXA - General", "regions": "DEXA - Enhanced"}
REQUIRED = {
    "manual": {"Date", "Weight", "Waist", "Body Fat %"},
    "evolt": {"Date", "Weight", "Lean Body Mass", "Total Body Fat Percentage"},
    "dexa": {"Date", "Total Mass", "Tissue % Fat"},
    "regions": {"Scan_Date", "Region", "Tissue_Fat_pct"},
}
PCT_FRACTION = {
    "evolt": {"Total Body Fat Percentage", "Waist To Hip Ratio"},
    "dexa": {"Tissue % Lean", "Tissue % Fat"},
}
# Original sheet and column labels are preserved in payloads; avoid mixing definitions.
class CompositionRow(db.Model):
    __tablename__ = "composition_rows"
    id = db.Column(db.Integer, primary_key=True)
    source = db.Column(db.String(20), nullable=False, index=True)
    observed_on = db.Column(db.Date, nullable=False, index=True)
    row_key = db.Column(db.String(100), nullable=False)
    payload = db.Column(db.JSON, nullable=False)
    __table_args__ = (db.UniqueConstraint("source", "row_key"),)

class CompositionImport(db.Model):
    __tablename__ = "composition_import"
    id = db.Column(db.Integer, primary_key=True)
    digest = db.Column(db.String(64), nullable=False, index=True)
    filename = db.Column(db.String(255), nullable=False)
    imported_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    counts = db.Column(db.JSON, nullable=False)
    warnings = db.Column(db.JSON, default=list)

def source_path():
    return Path(str(current_app.config.get("BODY_COMPOSITION_XLSX_PATH") or "")).expanduser()

def scalar(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    return str(value)[:2000]

def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None

def parse_workbook(path):
    """Validate all canonical worksheets before the database is touched."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        result, warnings = {}, []
        for kind, sheet_name in SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                raise ValueError("Required worksheet missing: " + sheet_name)
            ws = workbook[sheet_name]
            iterator = ws.iter_rows(values_only=True)
            header = [str(v).strip() if v is not None else "" for v in next(iterator)]
            if not REQUIRED[kind].issubset(set(header)):
                raise ValueError("Invalid columns in worksheet: " + sheet_name)
            records, keys = [], set()
            for row_number, values in enumerate(iterator, 2):
                raw = dict(zip(header, values))
                day = as_date(raw.get("Scan_Date" if kind == "regions" else "Date"))
                if day is None:
                    if any(v is not None for v in values):
                        warnings.append(sheet_name + ": invalid date at row " + str(row_number))
                    continue
                payload = {name: scalar(value) for name, value in raw.items() if name}
                for field in PCT_FRACTION.get(kind, set()):
                    value = payload.get(field)
                    if isinstance(value, (int, float)):
                        if not 0 <= value <= 1:
                            warnings.append(sheet_name + ": unusual fractional value at row " + str(row_number))
                        payload[field] = round(value * 100, 6) if "Percentage" in field or "%" in field else value
                if kind == "regions":
                    region = str(payload.get("Region") or "").strip()
                    if not region:
                        continue
                    key = day.isoformat() + "|" + region
                    for field in ("Tissue_g", "Fat_g", "Lean_g", "BMC_g"):
                        if isinstance(payload.get(field), (int, float)):
                            payload[field.replace("_g", "_kg")] = payload[field] / 1000
                elif kind == "evolt":
                    key = day.isoformat() + "|" + str(payload.get("Time") or "")
                else:
                    key = day.isoformat()
                if key in keys:
                    raise ValueError("Duplicate " + sheet_name + " key: " + key)
                keys.add(key)
                records.append((day, key, payload))
            result[kind] = records
        return result, warnings
    finally:
        workbook.close()

def import_source_bytes(blob, filename="Body Composition.xlsx", force=False):
    """Ingest validated remote bytes without requiring a synced local OneDrive file."""
    digest = hashlib.sha256(blob).hexdigest()
    previous = CompositionImport.query.order_by(CompositionImport.id.desc()).first()
    if previous and previous.digest == digest and not force:
        return {"status": "unchanged", "counts": previous.counts}
    records, warnings = parse_workbook(BytesIO(blob))
    counts = {kind: len(items) for kind, items in records.items()}
    try:
        CompositionRow.query.delete(synchronize_session=False)
        for kind, items in records.items():
            for day, key, payload in items:
                db.session.add(CompositionRow(source=kind, observed_on=day, row_key=key, payload=payload))
        db.session.add(CompositionImport(digest=digest, filename=Path(filename).name,
                                         counts=counts, warnings=warnings[:100]))
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {"status": "success", "counts": counts, "warnings": warnings[:100]}


def import_source(force=False):
    path = source_path()
    if not str(current_app.config.get("BODY_COMPOSITION_XLSX_PATH") or "").strip():
        return {"status": "unconfigured"}
    if not path.is_file():
        return {"status": "missing", "message": "Configured workbook unavailable"}
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    previous = CompositionImport.query.order_by(CompositionImport.id.desc()).first()
    if previous and previous.digest == digest and not force:
        return {"status": "unchanged", "counts": previous.counts}
    records, warnings = parse_workbook(path)
    counts = {kind: len(rows) for kind, rows in records.items()}
    try:
        CompositionRow.query.delete(synchronize_session=False)
        for kind, rows in records.items():
            for day, key, payload in rows:
                db.session.add(CompositionRow(source=kind, observed_on=day, row_key=key, payload=payload))
        db.session.add(CompositionImport(digest=digest, filename=path.name,
                                         counts=counts, warnings=warnings[:100]))
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {"status": "success", "counts": counts, "warnings": warnings[:100]}

def rows(kind, start=None, end=None):
    q = CompositionRow.query.filter_by(source=kind)
    if start:
        q = q.filter(CompositionRow.observed_on >= start)
    if end:
        q = q.filter(CompositionRow.observed_on <= end)
    return [{"date": r.observed_on.isoformat(), **r.payload} for r in
            q.order_by(CompositionRow.observed_on, CompositionRow.row_key).all()]

def latest(kind):
    r = CompositionRow.query.filter_by(source=kind).order_by(
        CompositionRow.observed_on.desc(), CompositionRow.id.desc()).first()
    return {"date": r.observed_on.isoformat(), **r.payload} if r else None

def overview_data():
    result = {kind: latest(kind) for kind in ("manual", "evolt", "dexa")}
    result["counts"] = {kind: CompositionRow.query.filter_by(source=kind).count() for kind in SHEETS}
    result["goals"] = {
        "weight_kg": float(current_app.config.get("BODY_WEIGHT_GOAL_KG") or 84),
        "body_fat_pct": float(current_app.config.get("BODY_FAT_GOAL_PCT") or 8)}
    return result

def status_data():
    last = CompositionImport.query.order_by(CompositionImport.id.desc()).first()
    path = source_path()
    configured = bool(str(current_app.config.get("BODY_COMPOSITION_XLSX_PATH") or "").strip())
    return {"status": "ready" if last else ("missing" if configured else "unconfigured"),
            "file_available": configured and path.is_file(),
            "last_import": last.imported_at.isoformat() if last else None,
            "counts": last.counts if last else {},
            "warnings": last.warnings if last else [],
            "source_filename": last.filename if last else None}

def filtered(kind):
    try:
        start = date.fromisoformat(request.args["from"]) if request.args.get("from") else None
        end = date.fromisoformat(request.args["to"]) if request.args.get("to") else None
        year = int(request.args["year"]) if request.args.get("year") else None
        month = int(request.args["month"]) if request.args.get("month") else None
        if month and not 1 <= month <= 12:
            raise ValueError()
        if year and not 1900 <= year <= 2200:
            raise ValueError()
        if month and not year:
            raise ValueError()
        if year:
            a = date(year, month or 1, 1)
            b = (date(year + (month == 12), (month % 12) + 1, 1) - timedelta(days=1)) if month else date(year, 12, 31)
            start = max(start, a) if start else a
            end = min(end, b) if end else b
        if start and end and start > end:
            raise ValueError()
    except (ValueError, OverflowError):
        raise ValueError("Invalid date filters")
    return rows(kind, start, end)

def projection(metric, horizon, window):
    if metric not in {"weight", "waist"} or horizon not in (30, 60, 90) or not 14 <= window <= 365:
        raise ValueError("Invalid projection parameters")
    field = "Weight" if metric == "weight" else "Waist"
    history = [(date.fromisoformat(r["date"]), r.get(field)) for r in rows("manual")]
    history = [(d, float(v)) for d, v in history if isinstance(v, (int, float)) and 0 < v < 500]
    if not history:
        return {"status": "insufficient_data", "points": []}
    anchor = history[-1][0]
    history = [(d, v) for d, v in history if (anchor - d).days <= window]
    if len(history) < 5 or (history[-1][0] - history[0][0]).days < 21:
        return {"status": "insufficient_data", "points": [], "sample_count": len(history)}
    xx = [(d-history[0][0]).days for d, _ in history]
    yy = [v for _, v in history]
    mx, my = sum(xx)/len(xx), sum(yy)/len(yy)
    denominator = sum((x-mx)**2 for x in xx)
    slope = sum((x-mx)*(y-my) for x, y in zip(xx, yy))/denominator if denominator else 0
    slope = max(-.15, min(.15, slope))
    spread = max(.5, math.sqrt(sum((y - (my+slope*(x-mx)))**2 for x,y in zip(xx,yy))/len(xx)))
    projected = []
    for offset in sorted(set([0, 30, 60, 90, horizon])):
        if offset > horizon:
            continue
        value = max(0, yy[-1] + slope*offset)
        band = spread * math.sqrt(1+offset/30)
        projected.append({"date": (anchor+timedelta(days=offset)).isoformat(),
                          "value": round(value,2), "lower": round(max(0,value-band),2),
                          "upper": round(value+band,2)})
    return {"status": "illustrative", "metric": metric, "source": "manual",
            "sample_count": len(history), "window_days": window, "horizon_days": horizon,
            "slope_per_week": round(slope*7,3), "points": projected,
            "note": "Capped linear extrapolation, not a medical prediction or calibrated confidence interval."}

@api.errorhandler(ValueError)
def invalid(e):
    return jsonify(error={"message": str(e)}), 400

@api.get("/overview")
def get_overview():
    return jsonify(data=overview_data(), meta={"sources": SHEETS})

@api.get("/measurements")
def get_manual():
    return jsonify(data=filtered("manual"), meta={"source": "Data", "units": "cm, kg, percentage points"})

@api.get("/evolt")
def get_evolt():
    return jsonify(data=filtered("evolt"), meta={"source": "Evolt Scans", "percentages": "percentage points"})

@api.get("/dexa/scans")
def get_dexa():
    return jsonify(data=filtered("dexa"), meta={"source": "DEXA - General", "percentages": "percentage points"})

@api.get("/dexa/regions")
def get_regions():
    data = filtered("regions")
    if request.args.get("date"):
        try:
            selected = date.fromisoformat(request.args["date"]).isoformat()
        except ValueError:
            raise ValueError("Invalid scan date")
        data = [r for r in data if r["date"] == selected]
    return jsonify(data=data, meta={"source": "DEXA - Enhanced", "mass": "grams and derived kilograms"})

@api.get("/projections")
def get_projection():
    try:
        horizon = int(request.args.get("horizon_days", 90))
        window = int(request.args.get("window_days", 90))
    except ValueError:
        raise ValueError("Invalid projection parameters")
    return jsonify(data=projection(request.args.get("metric", "weight"), horizon, window))

@api.get("/imports/status")
def get_status():
    return jsonify(data=status_data())

@pages.get("/fitness/body-composition")
@pages.get("/fitness/body-composition/measurements")
@pages.get("/fitness/body-composition/evolt")
@pages.get("/fitness/body-composition/dexa")
@pages.get("/fitness/body-composition/data")
def composition_page():
    section = request.path.rsplit("/", 1)[-1]
    if section == "body-composition":
        section = "overview"
    return render_template("body_composition.html", section="body-composition",
                           title="Body composition", view=section)
