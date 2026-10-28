"""Synthetic composition regression tests: no private workbook required."""
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from app import create_app, db
from app.body_composition import import_source, overview_data, projection, parse_workbook

def fixture_workbook(path):
    wb = Workbook()
    manual = wb.active
    manual.title = "Data"
    manual.append(["Date", "Weight", "Waist", "Body Fat %", "Chest", "Notes"])
    for day in (1, 8, 15, 22, 29, 30):
        manual.append([date(2026, 9, day), 90-day*.08, 88-day*.1, 11.79, 105, ""])
    ev = wb.create_sheet("Evolt Scans")
    ev.append(["Date ", "Time", "Weight", "Lean Body Mass", "Total Body Fat Percentage", "Abdomonal Circumference"])
    ev.append([date(2026, 6, 15), "08:00", 93.3, 84.2, .098, 84])
    dexa = wb.create_sheet("DEXA - General")
    dexa.append(["Date", "Total Mass", "Lean Mass", "Tissue % Fat", "Fat Mass"])
    dexa.append([date(2026, 9, 7), 90.4, 74.805, .142, 12.424])
    reg = wb.create_sheet("DEXA - Enhanced")
    reg.append(["Scan_Date", "Region", "Tissue_Fat_pct", "Fat_g", "Lean_g"])
    reg.append([date(2026, 9, 7), "Left Arm", 9.6, 300, 3000])
    reg.append([date(2026, 9, 7), "Right Arm", 9.5, 310, 3100])
    wb.save(path)

@pytest.fixture
def client(tmp_path):
    path=tmp_path/"body composition.xlsx"
    fixture_workbook(path)
    app=create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///"+str(tmp_path/"test.sqlite"),
                    "BODY_COMPOSITION_XLSX_PATH": str(path),
                    "BODY_COMPOSITION_IMPORT_ON_STARTUP": False, "WTF_CSRF_ENABLED": False})
    with app.app_context():
        yield app.test_client(),path
        db.session.remove()

def test_parse_percent_and_masses(client):
    _,path=client
    data,warnings=parse_workbook(path)
    assert len(data["manual"])==6 and len(data["regions"])==2
    assert data["manual"][-1][2]["Body Fat %"]==11.79
    assert data["evolt"][0][2]["Total Body Fat Percentage"]==9.8
    assert data["dexa"][0][2]["Tissue % Fat"]==14.2
    assert data["regions"][0][2]["Tissue_Fat_pct"]==9.6
    assert data["regions"][0][2]["Fat_kg"]==.3
    assert data["evolt"][0][2]["Abdomonal Circumference"]==84

def test_import_idempotent_and_endpoints(client):
    c,path=client
    with c.application.app_context():
        assert import_source()["status"]=="success"
        assert import_source()["status"]=="unchanged"
        assert overview_data()["counts"]["manual"]==6
    r=c.get("/api/v1/fitness/body-composition/overview")
    assert r.status_code==200 and r.json["data"]["evolt"]["Total Body Fat Percentage"]==9.8
    assert len(c.get("/api/v1/fitness/body-composition/dexa/regions").json["data"])==2
    assert c.get("/api/v1/fitness/body-composition/measurements?year=2026&month=13").status_code==400
    assert c.get("/fitness/body-composition").status_code==200

def test_changed_and_broken_source_preserves_existing(client):
    c,path=client
    with c.application.app_context():
        import_source()
        original=overview_data()["counts"]["manual"]
        path.write_bytes(b"not a workbook")
        with pytest.raises(Exception):
            import_source()
        assert overview_data()["counts"]["manual"]==original

def test_projections_are_bounded_and_source_specific(client):
    c,path=client
    with c.application.app_context():
        import_source()
        value=projection("weight",90,90)
        assert value["status"]=="illustrative"
        assert all(p["value"]>=0 for p in value["points"])
        with pytest.raises(ValueError):
            projection("dexa",90,90)
