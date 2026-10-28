"""Validate remote byte ingestion with synthetic private-data-free fixtures."""
from datetime import date
from io import BytesIO
import json
import zipfile

import pytest
from openpyxl import Workbook
from app import create_app
from app.body_composition import parse_workbook, import_source_bytes, CompositionRow
from app.gymday import load_zip, import_zip_bytes, GymSession


@pytest.fixture
def app(tmp_path):
    return create_app({"TESTING": True, "WTF_CSRF_ENABLED": False,
                       "SQLALCHEMY_DATABASE_URI": "sqlite:///" + str(tmp_path / "parsers.sqlite")})


def body_book():
    wb = Workbook()
    manual = wb.active
    manual.title = "Data"
    manual.append(["Date", "Weight", "Waist", "Body Fat %"])
    manual.append([date(2026, 9, 10), 88, 84, .15])
    evolt = wb.create_sheet("Evolt Scans")
    evolt.append(["Date", "Weight", "Lean Body Mass", "Total Body Fat Percentage"])
    evolt.append([date(2026, 9, 10), 88, 75, .15])
    dexa = wb.create_sheet("DEXA - General")
    dexa.append(["Date", "Total Mass", "Tissue % Fat"])
    dexa.append([date(2026, 9, 10), 88, .15])
    regions = wb.create_sheet("DEXA - Enhanced")
    regions.append(["Scan_Date", "Region", "Tissue_Fat_pct"])
    regions.append([date(2026, 9, 10), "Trunk", 15])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_body_source_bytes_preserve_local_schema(app, tmp_path):
    blob = body_book()
    path = tmp_path / "source.xlsx"
    path.write_bytes(blob)
    assert parse_workbook(path) == parse_workbook(BytesIO(blob))
    with app.app_context():
        assert import_source_bytes(blob)["status"] == "success"
        assert import_source_bytes(blob)["status"] == "unchanged"
        assert CompositionRow.query.count() == 4


def gym_zip(extra=None):
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("Exercise.json", json.dumps([{"id": "e1", "name": "Squat"}]))
        z.writestr("WorkoutSession.json", json.dumps([
            {"id": "s1", "name": "Leg Day", "startDate": 1790000000000}]))
        if extra:
            z.writestr(extra, "not allowed")
    return buf.getvalue()


def test_gym_remote_zip_bytes(app):
    payload = gym_zip()
    with app.app_context():
        assert import_zip_bytes(payload)["status"] == "success"
        assert import_zip_bytes(payload)["status"] == "unchanged"
        assert GymSession.query.count() == 1


def test_gym_zip_slip_rejected(app):
    with app.app_context():
        with pytest.raises(ValueError, match="Unsafe archive"):
            load_zip(BytesIO(gym_zip("../escape.txt")))


def test_gym_corrupt_zip_rejected(app):
    with app.app_context():
        with pytest.raises(zipfile.BadZipFile):
            import_zip_bytes(b"not a zip")
