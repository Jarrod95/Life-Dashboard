"""Regression contract for metadata skipping and failed network imports."""
from io import BytesIO
from openpyxl import Workbook
import pytest
from app import create_app
from app.models import DataSourceState, SyncRun
from app.services.data_sources import export_rank
from app.services.sync import sync_one


@pytest.fixture
def app(tmp_path):
    return create_app({"TESTING": True, "WTF_CSRF_ENABLED": False,
                       "LIFE_DATA_SOURCE": "onedrive", "LIFE_SYNC_LOCAL_FALLBACK": False,
                       "SQLALCHEMY_DATABASE_URI": "sqlite:///" + str(tmp_path / "test.sqlite")})


class FakeClient:
    def __init__(self, payload, etag="v1", fail=False):
        self.payload, self.etag, self.fail = payload, etag, fail
        self.downloads = 0
    def get_item_by_path(self, path):
        if self.fail:
            from app.integrations.onedrive import GraphError
            raise GraphError("unavailable", "Network error")
        return {"id": "remote-1", "name": "Financial Dashboard.xlsx",
                "parentReference": {"driveId": "d"}, "eTag": self.etag, "file": {}}
    def get_item_by_id(self, drive, item):
        return self.get_item_by_path("")
    def download(self, drive, item, max_bytes):
        self.downloads += 1
        return self.payload


def book():
    wb = Workbook()
    ws = wb.active
    ws.title = "Net Worth Tracking"
    ws["B1"] = "2026-01-01"
    ws["A12"] = "Cash"
    ws["B12"] = 100
    wb.create_sheet("Transactions_2026_NAB").append(
        ["Date", "Amount", "Account Number", "Transaction Type", "Transaction Details", "Category", "Merchant Name"])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def test_first_sync_and_etag_skip(app):
    with app.app_context():
        fake = FakeClient(book())
        first = sync_one("finance_workbook", client=fake)
        assert first["status"] == "success", first
        second = sync_one("finance_workbook", client=fake)
        assert second["status"] == "unchanged"
        assert fake.downloads == 1
        state = DataSourceState.query.filter_by(source_key="finance_workbook").one()
        assert state.sha256 and state.last_success_at
        assert SyncRun.query.count() == 2


def test_failure_retains_success_state(app):
    with app.app_context():
        fake = FakeClient(book())
        assert sync_one("finance_workbook", client=fake)["status"] == "success"
        last_hash = DataSourceState.query.first().sha256
        failed = sync_one("finance_workbook", client=FakeClient(book(), fail=True))
        assert failed["status"] == "failed"
        state = DataSourceState.query.first()
        assert state.sha256 == last_hash
        assert state.last_success_at is not None


def test_force_reprocess(app):
    with app.app_context():
        fake = FakeClient(book())
        assert sync_one("finance_workbook", client=fake)["status"] == "success"
        assert sync_one("finance_workbook", client=fake, force=True)["status"] == "success"
        assert fake.downloads == 2


def test_latest_export_order():
    a = {"name": "daily_strength_2026_10_01.zip", "id": "a",
         "lastModifiedDateTime": "2026-10-28T00:00:00Z"}
    b = {"name": "daily_strength_2026_10_21.zip", "id": "b",
         "lastModifiedDateTime": "2026-10-22T00:00:00Z"}
    assert max((a, b), key=export_rank) is b
