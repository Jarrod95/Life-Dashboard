"""Regression tests for source-backed mortgage summaries and Westpac payroll credits."""
from datetime import date
from decimal import Decimal
from openpyxl import Workbook
import pytest
from app import create_app


@pytest.fixture
def app(tmp_path):
    application = create_app({
        "TESTING": True,
        "WTF_CSRF_ENABLED": False,
        "SQLALCHEMY_DATABASE_URI": "sqlite:///" + str(tmp_path / "finance-test.sqlite"),
    })
    yield application

from app.finance import mortgage_tracking, spending
from app.importer import import_bytes


def test_mortgage_source_history_and_estimates(app, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Net Worth Tracking"
    ws["A54"] = "Mortgage balance"
    ws["A64"] = "Interest rate"
    ws["A65"] = "Offset balance"
    ws["A66"] = "Months remaining"
    for col, month, loan in [(2, 1, 400000), (3, 2, 398000), (4, 3, 396000)]:
        ws.cell(1, col, date(2026, month, 1))
        ws.cell(54, col, loan)
        ws.cell(64, col, 5.64)
        ws.cell(65, col, 60000)
        ws.cell(66, col, 320)
    path = tmp_path / "financial.xlsx"
    wb.save(path)
    with app.app_context():
        app.config["SOURCE_WORKBOOK"] = str(path)
        result = mortgage_tracking()
        assert result["status"] == "ok"
        assert len(result["history"]) == 3
        assert len(result["projection"]) == 12
        assert result["indicators"]["principal"] == "396000.00"
        assert result["indicators"]["interest_rate"] == str(Decimal("0.0564"))
        assert result["indicators"]["months_remaining"] == "320"
        assert Decimal(result["indicators"]["monthly_offset_interest_saving"]) > 0


def test_westpac_salary_credit_is_income(app):
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions_Westpac_All"
    ws.append(["Bank Account", "Date", "Narrative", "Debit Amount",
               "Credit Amount", "Balance", "Categories", "Serial"])
    ws.append(["A", date(2026, 3, 1), "DEPOSIT-SALARY Payroll",
               None, 4000, None, "DEP", None])
    ws.append(["A", date(2026, 3, 2), "Groceries",
               80, None, None, "Food", None])
    out = BytesIO()
    wb.save(out)
    with app.app_context():
        import_bytes(out.getvalue(), "financial.xlsx")
        rows = spending(group_by="month")
        march = next(row for row in rows if row["label"] == "2026-03")
        assert march["income"] == "4000.00"
        assert march["spending"] == "80.00"
