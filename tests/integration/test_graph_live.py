"""Explicit opt-in, read-only Microsoft Graph smoke test."""
import os
from io import BytesIO

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.graph_live


@pytest.mark.skipif(os.getenv("RUN_GRAPH_LIVE_TESTS") != "1",
                    reason="Set RUN_GRAPH_LIVE_TESTS=1 after user consent")
def test_graph_readonly_smoke():
    from app import create_app
    from app.integrations.onedrive import OneDriveClient
    from app.services.data_sources import SOURCES, resolve

    app = create_app({"TESTING": True})
    with app.app_context():
        client = OneDriveClient()
        assert client.whoami().get("id")
        assert client.drive().get("id")
        for key in SOURCES:
            item, drive = resolve(client, key)
            assert item["id"]
            print(key, "metadata OK", "type", "zip" if key == "gymday_latest_zip" else "xlsx")
            if key == "body_composition_workbook":
                data = client.download(drive, item["id"], max_bytes=30_000_000)
                wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
                try:
                    print("Body workbook sheets:", ", ".join(wb.sheetnames))
                finally:
                    wb.close()
